"""Dataset validation and language expansion for offline benchmarks."""

from __future__ import annotations

import csv
import hashlib
import json
import re
from pathlib import Path
from typing import Any, Iterator

LANGUAGES = ("english", "bangla", "banglish")
QUESTION_COLUMNS = {
    "english": ("english_question", "English Query", "english_query", "question_en"),
    "bangla": ("bangla_question", "Bengali Query", "bengali_query", "question_bn"),
    "banglish": ("banglish_question", "Banglish Query", "banglish_query", "question_bl"),
}
REFERENCE_COLUMNS = {
    "english": ("ground_truth_answer_en", "reference_answer_en", "reference_answer"),
    "bangla": ("ground_truth_answer_bn", "reference_answer_bn", "reference_answer"),
    "banglish": ("ground_truth_answer_banglish", "reference_answer_banglish", "reference_answer"),
}
VALID_ANNOTATIONS = {"verified", "needs_review", "source_conflict", "ambiguous_reference", "unsupported"}
_SPLIT_LABELS = re.compile(r"\s*[;|]\s*")


def first(row: dict[str, Any], names: tuple[str, ...] | list[str]) -> str:
    for name in names:
        value = row.get(name)
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""


def evidence_labels(value: Any, *, pages: bool = False) -> list[str]:
    """Parse single, semicolon/pipe-separated, or JSON-array evidence labels."""
    if value is None or str(value).strip() == "":
        return []
    if pages and isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (list, tuple)):
        parts = [str(item).strip() for item in value]
    else:
        raw = str(value).strip()
        if raw.startswith("["):
            parsed = json.loads(raw)
            if not isinstance(parsed, list):
                raise ValueError("evidence label JSON must be an array")
            parts = [str(item).strip() for item in parsed]
        else:
            parts = _SPLIT_LABELS.split(raw)
    if any(not item for item in parts):
        raise ValueError("evidence labels may not contain empty entries")
    if pages:
        for item in parts:
            if not re.fullmatch(r"[1-9]\d*", item):
                raise ValueError(f"invalid expected page: {item!r}")
    return list(dict.fromkeys(parts))


def _read_csv(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    with path.open("r", encoding="utf-8-sig", errors="strict", newline="") as handle:
        reader = csv.DictReader(handle)
        columns = reader.fieldnames or []
        if not columns or len(columns) != len(set(columns)):
            raise ValueError("CSV header is missing or contains duplicate columns")
        rows = list(reader)
    if any(None in row for row in rows):
        raise ValueError("CSV row has more values than its header")
    return columns, rows


def _read_xlsx(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        header = next(iterator, None)
        columns = [str(value).strip() if value is not None else "" for value in (header or ())]
        if not columns or any(not col for col in columns) or len(columns) != len(set(columns)):
            raise ValueError("XLSX header is missing, blank, or contains duplicate columns")
        rows = []
        for cells in iterator:
            if all(value is None or str(value).strip() == "" for value in cells):
                continue
            if any(value is not None for value in cells[len(columns):]):
                raise ValueError("XLSX row has values beyond the header")
            rows.append(dict(zip(columns, cells)))
        return columns, rows
    finally:
        workbook.close()


def load_and_validate(path: str | Path, *, languages: tuple[str, ...] = LANGUAGES,
                      require_references: bool = False) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    path = Path(path).resolve()
    if path.suffix.casefold() == ".csv":
        columns, rows = _read_csv(path)
    elif path.suffix.casefold() == ".xlsx":
        columns, rows = _read_xlsx(path)
    else:
        raise ValueError("dataset must be CSV or XLSX")
    if not rows:
        raise ValueError("dataset has no rows")
    unknown_languages = set(languages) - set(LANGUAGES)
    if not languages or unknown_languages:
        raise ValueError(f"invalid language selection: {sorted(unknown_languages)}")
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    errors: list[str] = []
    warnings: list[str] = []
    seen_ids: set[str] = set()
    variants = 0
    normalized = []
    for number, original in enumerate(rows, 1):
        row = {str(key): value for key, value in original.items()}
        qid = first(row, ("question_id", "Question ID", "ID", "id")) or f"row-{number:06d}"
        if qid in seen_ids:
            errors.append(f"row {number}: duplicate question_id {qid!r}")
        seen_ids.add(qid)
        annotation = first(row, ("ground_truth_status", "annotation_status")) or "unlabeled"
        if annotation != "unlabeled" and annotation not in VALID_ANNOTATIONS:
            errors.append(f"row {number}: invalid ground_truth_status {annotation!r}")
        try:
            sources = evidence_labels(first(row, ("expected_sources", "expected_source")))
            pages = evidence_labels(first(row, ("expected_pages", "expected_page", "expected_page_of_answer")), pages=True)
        except (ValueError, json.JSONDecodeError) as exc:
            errors.append(f"row {number}: {exc}")
            sources, pages = [], []
        pairs = []
        raw_pairs = first(row, ("expected_evidence",))
        if raw_pairs:
            try:
                parsed = json.loads(raw_pairs)
                if not isinstance(parsed, list):
                    raise ValueError("expected_evidence must be a JSON array")
                for item in parsed:
                    if not isinstance(item, dict):
                        raise ValueError("expected_evidence entries must be objects")
                    pair_pages = evidence_labels(item.get("page"), pages=True)
                    pair_sources = evidence_labels(item.get("source"))
                    if len(pair_pages) != 1 or len(pair_sources) != 1:
                        raise ValueError("each expected_evidence entry requires one source and one page")
                    pairs.append({"source": pair_sources[0], "page": pair_pages[0]})
            except (ValueError, json.JSONDecodeError) as exc:
                errors.append(f"row {number}: {exc}")
        row["_base_question_id"] = qid
        row["_annotation_status"] = annotation
        row["_expected_sources"] = sources
        row["_expected_pages"] = pages
        row["_expected_evidence"] = pairs
        present = [language for language in languages if first(row, QUESTION_COLUMNS[language])]
        if not present:
            errors.append(f"row {number}: no question text for selected languages")
        for language in present:
            variants += 1
            if require_references and not first(row, REFERENCE_COLUMNS[language]):
                errors.append(f"row {number}: missing {language} reference answer")
            elif not first(row, REFERENCE_COLUMNS[language]):
                warnings.append(f"row {number}: {language} reference unavailable; answer proxies will be NOT_AVAILABLE")
        normalized.append(row)
    validation = {
        "valid": not errors, "errors": errors, "warnings": warnings,
        "dataset_filename": path.name, "dataset_path": str(path), "dataset_sha256": digest,
        "row_count": len(rows), "language_variant_count": variants,
        "column_schema": columns, "languages": list(languages),
    }
    return normalized, validation


def expand(rows: list[dict[str, Any]], languages: tuple[str, ...] = LANGUAGES) -> Iterator[dict[str, Any]]:
    seen: set[str] = set()
    for row in rows:
        base_id = row["_base_question_id"]
        for language in languages:
            question = first(row, QUESTION_COLUMNS[language])
            if not question:
                continue
            evaluation_id = f"{base_id}:{language}"
            if evaluation_id in seen:
                raise ValueError(f"duplicate evaluation_id {evaluation_id}")
            seen.add(evaluation_id)
            yield {
                "base_question_id": base_id, "evaluation_id": evaluation_id,
                "language": language, "question": question,
                "reference_answer": first(row, REFERENCE_COLUMNS[language]),
                "expected_sources": list(row["_expected_sources"]),
                "expected_pages": list(row["_expected_pages"]),
                "expected_evidence": list(row["_expected_evidence"]),
                "annotation_status": row["_annotation_status"],
                "intent": first(row, ("intent", "expected_intent")),
                "difficulty": first(row, ("difficulty",)),
                "course_code": first(row, ("course_code",)),
                "notes": first(row, ("notes",)),
            }
